# -*- coding: utf-8 -*-
"""Draw traffic graphs/charts and insert them into the AU case-study DOCX."""
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Circle, FancyArrowPatch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

BASE = Path(r"C:\Users\NITIN\Downloads\BloodConnect360_Frontend")
FIG = BASE / "_figures"
FIG.mkdir(exist_ok=True)
SRC = BASE / "Case_Study_Network_Traffic_Analysis_AU_Format.docx"
DOCX = BASE / "Case_Study_Network_Traffic_Analysis_AU_Format_With_Graphs.docx"

NAVY = "#1F4E79"
RED = "#C0392B"
ORANGE = "#E67E22"
GREY = "#7F8C8D"
PURPLE = "#6C3483"
TEAL = "#148F77"


def node(ax, xy, label, color, r=0.18, fs=8, tc="white"):
    ax.add_patch(Circle(xy, r, facecolor=color, edgecolor="white", lw=1.2, zorder=3))
    ax.text(xy[0], xy[1], label, ha="center", va="center", fontsize=fs,
            color=tc, fontweight="bold", zorder=4)


def edge(ax, a, b, color="#95A5A6", lw=1.2, rad=0):
    ax.add_patch(
        FancyArrowPatch(
            a, b, arrowstyle="-|>", mutation_scale=9, color=color, lw=lw,
            connectionstyle=f"arc3,rad={rad}", shrinkA=16, shrinkB=16, zorder=2,
        )
    )


def style_ax(ax, title):
    ax.set_xlim(-1.35, 1.35)
    ax.set_ylim(-1.45, 1.35)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title(title, fontsize=10, fontweight="bold", color=NAVY, pad=4)


def fig_topologies():
    fig, axes = plt.subplots(1, 4, figsize=(11.2, 3.15), dpi=160)
    fig.patch.set_facecolor("white")

    # Normal: two servers + clients
    ax = axes[0]
    servers = [(-0.15, 0.35), (0.35, -0.15)]
    clients = [(-0.9, 0.85), (-0.95, -0.15), (-0.55, -0.95), (0.15, -1.05),
               (0.95, -0.75), (1.05, 0.25), (0.7, 0.95), (-0.2, 1.05)]
    for c in clients:
        edge(ax, c, servers[0] if c[1] > 0 else servers[1], "#B0BEC5", 1.0)
    edge(ax, servers[0], servers[1], NAVY, 1.6)
    for i, c in enumerate(clients):
        node(ax, c, "", GREY, 0.10)
    node(ax, servers[0], "S1", NAVY, 0.20, 7)
    node(ax, servers[1], "S2", NAVY, 0.20, 7)
    style_ax(ax, "Normal traffic")
    ax.text(0, -1.38, "Few hubs, low client degree", ha="center", fontsize=7.5, color="#555")

    # DDoS star into victim
    ax = axes[1]
    import math
    victim = (0, 0)
    attackers = [(math.cos(a) * 1.05, math.sin(a) * 1.05) for a in
                 [i * math.pi / 4 for i in range(8)]]
    for a in attackers:
        edge(ax, a, victim, RED, 1.15)
        node(ax, a, "", RED, 0.11)
    node(ax, victim, "V", ORANGE, 0.24, 9)
    style_ax(ax, "DDoS (star in)")
    ax.text(0, -1.38, "High in-degree at victim", ha="center", fontsize=7.5, color="#555")

    # Scanner star out
    ax = axes[2]
    scanner = (-0.85, 0)
    targets = [(0.55, y) for y in (0.95, 0.48, 0.0, -0.48, -0.95)]
    for t in targets:
        edge(ax, scanner, t, RED, 1.15)
        node(ax, t, "", GREY, 0.12)
    node(ax, scanner, "Sc", RED, 0.24, 8)
    style_ax(ax, "Scan (star out)")
    ax.text(0, -1.38, "High out-degree at scanner", ha="center", fontsize=7.5, color="#555")

    # Botnet hub
    ax = axes[3]
    c2 = (0, 0)
    bots = [(math.cos(a) * 1.0, math.sin(a) * 1.0) for a in
            [i * 2 * math.pi / 7 - 0.2 for i in range(7)]]
    for b in bots:
        edge(ax, b, c2, PURPLE, 1.1)
        node(ax, b, "", PURPLE, 0.11)
    node(ax, c2, "C2", NAVY, 0.24, 8)
    style_ax(ax, "Botnet C2")
    ax.text(0, -1.38, "Hub-and-spoke community", ha="center", fontsize=7.5, color="#555")

    fig.tight_layout(pad=0.4)
    path = FIG / "fig1_topologies.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def fig_degree_timeline():
    times = ["14:10", "14:11", "14:12", "14:13", "14:14", "14:15", "14:16"]
    degree = [4, 5, 48, 420, 900, 860, 12]
    fig, ax = plt.subplots(figsize=(7.4, 3.15), dpi=160)
    fig.patch.set_facecolor("white")
    ax.plot(times, degree, color=RED, marker="o", lw=2.2, ms=6, label="IoT camera out-degree")
    ax.axhline(40, color=NAVY, ls="--", lw=1.3, label="Alert threshold (baseline + margin)")
    ax.axvspan(2.5, 5.5, color="#F5B7B1", alpha=0.45, label="Scan windows")
    ax.set_ylabel("Out-degree (distinct peers)", fontsize=9)
    ax.set_xlabel("One-minute windows", fontsize=9)
    ax.set_title("Figure context: campus IoT scan — structural spike, not a byte spike",
                 fontsize=10, color=NAVY, fontweight="bold")
    ax.set_ylim(0, 1050)
    ax.grid(axis="y", color="#E5E8E8", lw=0.8)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(frameon=False, fontsize=8, loc="upper left")
    ax.annotate("Isolate host", xy=(6, 12), xytext=(5.05, 280),
                fontsize=8, color=TEAL,
                arrowprops=dict(arrowstyle="->", color=TEAL, lw=1))
    fig.tight_layout()
    path = FIG / "fig2_degree_timeline.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def fig_feature_bars():
    labels = ["Max\nout-degree", "New edges", "Entropy\n(scaled ×10)", "Max\nin-degree"]
    normal = [18, 22, 31, 40]
    attack = [900, 860, 18, 46]
    import numpy as np
    x = np.arange(len(labels))
    w = 0.36
    fig, ax = plt.subplots(figsize=(7.2, 3.15), dpi=160)
    fig.patch.set_facecolor("white")
    ax.bar(x - w / 2, normal, w, color=NAVY, label="Normal window")
    ax.bar(x + w / 2, attack, w, color=RED, label="Scan window (14:14)")
    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=8.5)
    ax.set_ylabel("Feature value", fontsize=9)
    ax.set_title("Normal window vs scan window (same host segment)",
                 fontsize=10, color=NAVY, fontweight="bold")
    ax.legend(frameon=False, fontsize=8)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(axis="y", color="#E5E8E8", lw=0.8)
    fig.tight_layout()
    path = FIG / "fig3_feature_compare.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def insert_after(paragraph, image_path, width_in, caption):
    cap_el = OxmlElement("w:p")
    paragraph._p.addnext(cap_el)
    cap = Paragraph(cap_el, paragraph._parent)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(8)
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(8)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    img_el = OxmlElement("w:p")
    paragraph._p.addnext(img_el)
    img_p = Paragraph(img_el, paragraph._parent)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.paragraph_format.space_before = Pt(6)
    img_p.paragraph_format.space_after = Pt(2)
    img_p.add_run().add_picture(str(image_path), width=Inches(width_in))


def find_para(doc, startswith):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise SystemExit(f"Paragraph not found: {startswith!r}")


def main():
    f1 = fig_topologies()
    f2 = fig_degree_timeline()
    f3 = fig_feature_bars()
    doc = Document(str(SRC))

    insert_after(
        find_para(doc, "Finding. The hardest cases"),
        f1, 6.4,
        "Figure 1. Traffic-graph shapes: normal communication, DDoS star, scanner star and botnet hub-and-spoke.",
    )
    insert_after(
        find_para(doc, "Cross-check with the mini-scenario"),
        f2, 5.6,
        "Figure 2. Out-degree of the compromised IoT camera across one-minute windows (Table 9 scenario).",
    )
    insert_after(
        find_para(doc, "Temporal views also support situational awareness"),
        f3, 5.5,
        "Figure 3. Structural features in a normal window versus the scan window — out-degree and new edges dominate.",
    )

    doc.save(str(DOCX))
    print("Saved", DOCX)
    print("Figures:", f1.name, f2.name, f3.name)


if __name__ == "__main__":
    main()
